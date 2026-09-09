"""Renders a format-agnostic ReportContent into PDF/XLSX/CSV bytes.
architecture §4: "PDF (ReportLab/WeasyPrint), XLSX (openpyxl), CSV
(native)" — ReportLab over WeasyPrint since it's pure-Python (no
system Cairo/Pango dependency), matching this codebase's preference for
adapters that need no external service to implement and verify for
real (same reasoning as LocalFilesystemStorage over a cloud SDK)."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.reports.content import ReportContent


def render_pdf(content: ReportContent) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=content.title)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(content.title, styles["Title"]),
        Paragraph(
            f"{content.organisation_name} — {content.scope_description} — generated "
            f"{content.generated_at.strftime('%Y-%m-%d %H:%M')} UTC",
            styles["Normal"],
        ),
        Spacer(1, 16),
    ]
    for section in content.sections:
        story.append(Paragraph(section.title, styles["Heading2"]))
        if section.summary_fields:
            data = [[label, value] for label, value in section.summary_fields]
            table = Table(data, colWidths=[220, 220])
            table.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            story.append(table)
        if section.is_table:
            if not section.rows:
                story.append(Paragraph("No records.", styles["Normal"]))
            else:
                data = [section.columns, *section.rows]
                table = Table(data, repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                            ("FONTSIZE", (0, 0), (-1, -1), 8),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                            ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ]
                    )
                )
                story.append(table)
        story.append(Spacer(1, 14))
    doc.build(story)
    return buffer.getvalue()


def render_xlsx(content: ReportContent) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = content.title[:31] or "Report"
    bold = Font(bold=True)

    sheet.append([content.title])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"{content.organisation_name} — {content.scope_description}"])
    sheet.append([f"Generated {content.generated_at.strftime('%Y-%m-%d %H:%M')} UTC"])
    sheet.append([])

    for section in content.sections:
        sheet.append([section.title])
        sheet.cell(row=sheet.max_row, column=1).font = bold
        if section.summary_fields:
            for label, value in section.summary_fields:
                sheet.append([label, value])
        if section.is_table:
            header_row = sheet.max_row + 1
            sheet.append(section.columns)
            for cell in sheet[header_row]:
                cell.font = bold
            for row in section.rows:
                sheet.append(row)
        sheet.append([])

    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_csv(content: ReportContent) -> bytes:
    """CSV is inherently flat: headline fields render as key,value rows,
    then the report's primary table (its first table section — see each
    content.py builder's `csv_note` for what that means per report
    type), never every section concatenated into one ambiguous sheet."""

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([content.title])
    writer.writerow([content.organisation_name, content.scope_description])
    writer.writerow([f"Generated {content.generated_at.strftime('%Y-%m-%d %H:%M')} UTC"])
    writer.writerow([])

    for section in content.sections:
        if section.summary_fields:
            writer.writerow([section.title])
            for label, value in section.summary_fields:
                writer.writerow([label, value])
            writer.writerow([])

    primary_table = next((s for s in content.sections if s.is_table), None)
    if primary_table is not None:
        writer.writerow([primary_table.title])
        writer.writerow(primary_table.columns)
        writer.writerows(primary_table.rows)

    return buffer.getvalue().encode("utf-8")


RENDERERS = {"PDF": render_pdf, "XLSX": render_xlsx, "CSV": render_csv}
CONTENT_TYPES = {
    "PDF": "application/pdf",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "CSV": "text/csv",
}
EXTENSIONS = {"PDF": "pdf", "XLSX": "xlsx", "CSV": "csv"}
